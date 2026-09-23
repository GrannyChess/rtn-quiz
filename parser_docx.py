#!/usr/bin/env python3
"""
parser_docx.py v5 — определяет правильные ответы по rStyle,
которым Word пометил их при экспорте из HTML/Markdown.

Плюс оставлен fallback на прямое форматирование (bold/italic/
underline/highlight/color), если где-то оно есть.
"""
import json
import re
import sys
from collections import Counter
from pathlib import Path
from docx import Document

RE_QUESTION = re.compile(r'^\s*(\d{1,4})\.\s+(.+?)\s*$')
RE_OPTION   = re.compile(r'^\s*([А-Еа-еA-E])\)\s*(.+?)\s*$')
RE_MD_LINK  = re.compile(r'\[([^\]]+)\]\([^)]+\)')
RE_PAGE     = re.compile(r'^=+\s*Page\s*\d+\s*=+$', re.I)
RE_NOISE    = re.compile(
    r'(Актуально на|Материал из Справочной|^https?://|'
    r'Дата копирования|^Б\.12\.1\.\s*Вопросы)', re.I)

RE_RSTYLE = re.compile(r'<w:rStyle\s+w:val="([^"]+)"')
RE_BOLD   = re.compile(r'<w:b(?:\s+w:val="([^"]*)")?\s*/?>')
RE_ITAL   = re.compile(r'<w:i(?:\s+w:val="([^"]*)")?\s*/?>')
RE_UNDER  = re.compile(r'<w:u(?:\s+w:val="([^"]*)")?\s*/?>')
RE_HL     = re.compile(r'<w:highlight\s+w:val="([^"]+)"')
RE_COLOR  = re.compile(r'<w:color\s+w:val="([^"]+)"')


def _val_true(v):
    if v is None:
        return True
    return str(v).lower() not in ('0', 'false', 'off', 'none')


def run_rstyle(run):
    m = RE_RSTYLE.search(run._r.xml)
    return m.group(1) if m else None


def run_marked_direct(run):
    xml = run._r.xml
    for rx in (RE_BOLD, RE_ITAL, RE_UNDER):
        m = rx.search(xml)
        if m and _val_true(m.group(1)):
            return True
    if RE_HL.search(xml):
        return True
    m = RE_COLOR.search(xml)
    if m and m.group(1).upper() not in ('000000', 'FFFFFF', 'AUTO'):
        return True
    return False


def discover_correct_rstyles(doc):
    """
    Смотрит на абзацы-варианты (А), Б), ...).
    Считает, какие rStyle встречаются у них.
    Кандидаты в «правильный ответ» — те, что встречаются
    у 5%–95% вариантов: не у всех (иначе дефолт) и не у 1-2
    (иначе шум).
    """
    counts, total = Counter(), 0
    for p in doc.paragraphs:
        if RE_OPTION.match(p.text.strip()):
            total += 1
            styles = {run_rstyle(r) for r in p.runs if run_rstyle(r)}
            for s in styles:
                counts[s] += 1

    print(f"\nВсего абзацев-вариантов: {total}")
    print("rStyle в вариантах:")
    for s, c in counts.most_common():
        pct = c / total * 100 if total else 0
        tag = ' ★' if 5 <= pct <= 95 else ''
        print(f"  {s}: {c} ({pct:.1f}%){tag}")

    return {s for s, c in counts.items()
            if total and 5 <= c / total * 100 <= 95}


def para_is_marked(p, correct_styles):
    for r in p.runs:
        if run_marked_direct(r):
            return True
        s = run_rstyle(r)
        if s and s in correct_styles:
            return True
    return False


def iter_blocks(doc, correct_styles):
    for p in doc.paragraphs:
        yield p.text, para_is_marked(p, correct_styles)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p.text, para_is_marked(p, correct_styles)


def clean_ref(text):
    text = RE_MD_LINK.sub(r'\1', text)
    return re.sub(r'\s+', ' ', text).strip()


def parse_docx(path):
    doc = Document(path)
    correct_styles = discover_correct_rstyles(doc)
    print(f"\nИспользуются как маркеры: {sorted(correct_styles) or '— нет —'}")

    questions, seen = [], set()
    current, phase = None, None

    def flush():
        nonlocal current, phase
        if not current:
            return
        q = current
        q['question'] = re.sub(r'\s+', ' ', q['question']).strip()
        q['options'] = [re.sub(r'\s+', ' ', o).strip() for o in q['options']]
        q['reference'] = clean_ref(' '.join(q.pop('_ref')))
        if len(q['correct']) > 1:
            q['multi'] = True
        m = re.search(r'Выберите\s+(\d+)\s+вариант', q['question'], re.I)
        if m and int(m.group(1)) > 1:
            q['multi'] = True
        probs = []
        if not q['options']:
            probs.append('нет вариантов')
        if not q['correct']:
            probs.append('нет ответа')
        if probs:
            print(f"  ⚠ id={q['id']:>3}  {', '.join(probs)}  | {q['question'][:70]}…")
        # Фильтр: выбрасываем блоки без вариантов —
        # это случайно захваченный текст НПА, а не вопрос
        if not q['options']:
            print(f"  ⓘ пропущен блок id={q['id']} (без вариантов): "
                  f"{q['question'][:60]}…")
            current, phase = None, None
            return
        questions.append(q)
        current, phase = None, None

    def new_q(qid, text):
        nonlocal current, phase
        flush()
        if qid in seen:
            return
        seen.add(qid)
        current = {'id': qid, 'question': text,
                   'options': [], 'correct': [], '_ref': []}
        phase = 'question'

    for text, marked in iter_blocks(doc, correct_styles):
        lines = [x.strip() for x in re.split(r'[\n\r\v]+', text) if x.strip()]
        for line in lines:
            line = line.replace('\xa0', ' ').strip()
            if not line or RE_PAGE.match(line):
                continue
            if RE_NOISE.search(line) and not RE_QUESTION.match(line):
                continue

            m_q = RE_QUESTION.match(line)
            if m_q and phase != 'options':
                qid = int(m_q.group(1))
                is_next = (current is None and
                           (not questions or
                            qid == max(x['id'] for x in questions) + 1))
                if marked or is_next or current is None or current['options']:
                    new_q(qid, m_q.group(2))
                    continue

            if current is None:
                continue

            m_o = RE_OPTION.match(line)
            if m_o and phase in ('question', 'options'):
                letters = 'АБВГДЕ'
                idx = len(current['options'])
                expected = letters[idx] if idx < len(letters) else None
                letter = m_o.group(1).upper()
                letter = {'A':'А','B':'В','C':'С','D':'Д','E':'Е'}.get(letter, letter)
                if expected is None or letter == expected or (idx == 0 and letter == 'А'):
                    current['options'].append(m_o.group(2).strip())
                    if marked:
                        current['correct'].append(idx)
                    phase = 'options'
                    continue

            if phase == 'question':
                current['question'] += ' ' + line
                continue
            if phase in ('options', 'reference'):
                current['_ref'].append(line)
                phase = 'reference'

    flush()
    questions.sort(key=lambda q: q['id'])
    return questions


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        print("\nПример:")
        print('  python parser_docx.py b12.docx "Б.12.1" '
              '"Взрывные работы..." tests/b12_1.json')
        sys.exit(1)

    src = Path(sys.argv[1])
    code = sys.argv[2] if len(sys.argv) > 2 else src.stem
    title = sys.argv[3] if len(sys.argv) > 3 else ""
    dst = Path(sys.argv[4]) if len(sys.argv) > 4 else Path("tests") / (src.stem + ".json")

    if not src.exists():
        print(f"Файл не найден: {src}")
        sys.exit(1)

    dst.parent.mkdir(parents=True, exist_ok=True)

    print(f"Читаю: {src}")
    qs = parse_docx(str(src))

    single = sum(1 for q in qs if not q.get('multi'))
    multi  = sum(1 for q in qs if q.get('multi'))
    no_ans = [q['id'] for q in qs if not q['correct']]
    print(f"\nРаспознано: {len(qs)}")
    print(f"  один ответ:   {single}")
    print(f"  мультивыбор:  {multi}")
    if no_ans:
        print(f"  ⚠ без ответа:  {no_ans}")

    payload = {
        "_meta": {
            "code": code,
            "title": title,
            "source": "",
        },
        "questions": qs,
    }
    with dst.open('w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(f"\nЗаписано: {dst}")


if __name__ == '__main__':
    main()