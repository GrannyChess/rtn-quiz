# peek_ids.py — смотрит, как помечены варианты у конкретных вопросов
import re
import sys
from collections import Counter
from docx import Document

path = sys.argv[1]
want_ids = [int(x) for x in sys.argv[2:]]

doc = Document(path)

RE_Q     = re.compile(r'^\s*(\d{1,4})\.\s+(.+?)\s*$')
RE_OPT   = re.compile(r'^\s*([А-Еа-еA-E])\)\s*(.+?)\s*$')
RE_RSTYLE = re.compile(r'<w:rStyle\s+w:val="([^"]+)"')
RE_BOLD  = re.compile(r'<w:b(?:\s+w:val="([^"]*)")?\s*/?>')
RE_ITAL  = re.compile(r'<w:i(?:\s+w:val="([^"]*)")?\s*/?>')
RE_UNDER = re.compile(r'<w:u(?:\s+w:val="([^"]*)")?\s*/?>')
RE_COLOR = re.compile(r'<w:color\s+w:val="([^"]+)"')
RE_HL    = re.compile(r'<w:highlight\s+w:val="([^"]+)"')


def marks(run):
    xml = run._r.xml
    out = []
    m = RE_RSTYLE.search(xml)
    if m: out.append(f"rStyle={m.group(1)}")
    m = RE_BOLD.search(xml)
    if m: out.append("bold")
    m = RE_ITAL.search(xml)
    if m: out.append("italic")
    m = RE_UNDER.search(xml)
    if m: out.append("underline")
    m = RE_COLOR.search(xml)
    if m and m.group(1).upper() not in ('000000', 'FFFFFF', 'AUTO'):
        out.append(f"color={m.group(1)}")
    if RE_HL.search(xml):
        out.append("highlight")
    return out or ["—"]


# 1) Все rStyle в документе с частотой
all_r = Counter()
for p in doc.paragraphs:
    for r in p.runs:
        m = RE_RSTYLE.search(r._r.xml)
        if m:
            all_r[m.group(1)] += 1

print("=== Все rStyle в документе (частота по runs) ===")
for k, v in all_r.most_common():
    print(f"  {k}: {v}")

# 2) Подробно по каждому запрошенному id
print(f"\n=== Детали вопросов: {want_ids} ===")
current_id = None
phase = None

for p in doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue

    m = RE_Q.match(t)
    if m and phase != 'options':
        current_id = int(m.group(1))
        phase = 'q'
        if current_id in want_ids:
            print(f"\n──── Q{current_id}: {m.group(2)[:90]}")
        continue

    mo = RE_OPT.match(t)
    if mo and current_id in want_ids and phase in ('q', 'options'):
        letter = mo.group(1).upper()
        text = mo.group(2)[:90]
        print(f"  {letter}) {text}")
        for i, r in enumerate(p.runs):
            mr = marks(r)
            print(f"     run[{i}] {r.text[:50]!r}  →  {', '.join(mr)}")
        phase = 'options'
        continue

    if current_id in want_ids and phase == 'options':
        if RE_OPT.match(t):
            continue
        # ссылка/текст после вариантов
        phase = 'ref'